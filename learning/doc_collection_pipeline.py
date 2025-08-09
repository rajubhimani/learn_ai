import os
import json
import boto3
from pathlib import Path
from typing import List, Dict, Any, Optional
from dataclasses import dataclass
from datetime import datetime

import pandas as pd
import chromadb
from chromadb.config import Settings
from chromadb.errors import NotFoundError
import qdrant_client
from qdrant_client.http.models import Distance, VectorParams
from sentence_transformers import SentenceTransformer
import hashlib
from chromadb.api.types import QueryResult, Metadata

# Document parsers
from pypdf import PdfReader
import docx
from bs4 import BeautifulSoup
import csv

@dataclass
class Document:
    content: str
    metadata: Dict[str, Any]
    doc_id: str
    chunk_id: Optional[str] = None

class FileProcessor:
    """Handles different file formats and sources"""
    
    def __init__(self):
        self.s3_client = boto3.client('s3') if self._has_aws_credentials() else None
        
    def _has_aws_credentials(self) -> bool:
        """Check if AWS credentials are available"""
        try:
            boto3.Session().get_credentials()
            return True
        except:
            return False
    
    def process_file(self, file_path: str, source_type: str = "local") -> List[Document]:
        """Process a file from local storage or S3"""
        if source_type == "s3":
            return self._process_s3_file(file_path)
        else:
            return self._process_local_file(file_path)
    
    def _process_local_file(self, file_path: str) -> List[Document]:
        """Process local file"""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_extension = path.suffix.lower()
        
        with open(file_path, 'rb') as f:
            content = f.read()
        
        return self._parse_content(content, file_extension, str(path))
    
    def _process_s3_file(self, s3_path: str) -> List[Document]:
        """Process S3 file (format: bucket/key/path)"""
        if not self.s3_client:
            raise ValueError("AWS credentials not configured")
        
        parts = s3_path.split('/', 1)
        bucket = parts[0]
        key = parts[1]
        
        response = self.s3_client.get_object(Bucket=bucket, Key=key)
        content = response['Body'].read()
        
        file_extension = Path(key).suffix.lower()
        return self._parse_content(content, file_extension, s3_path)
    
    def _parse_content(self, content: bytes, file_extension: str, source_path: str) -> List[Document]:
        """Parse content based on file extension"""
        doc_id = hashlib.md5(source_path.encode()).hexdigest()
        base_metadata = {
            'source': source_path,
            'file_type': file_extension,
            'processed_at': datetime.now().isoformat(),
            'doc_id': doc_id
        }
        
        if file_extension == '.pdf':
            return self._parse_pdf(content, base_metadata)
        elif file_extension == '.docx':
            return self._parse_docx(content, base_metadata)
        elif file_extension == '.txt':
            return self._parse_text(content, base_metadata)
        elif file_extension == '.json':
            return self._parse_json(content, base_metadata)
        elif file_extension == '.csv':
            return self._parse_csv(content, base_metadata)
        elif file_extension == '.html':
            return self._parse_html(content, base_metadata)
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")
    
    def _parse_pdf(self, content: bytes, metadata: Dict) -> List[Document]:
        """Parse PDF content"""
        import io
        pdf_file = io.BytesIO(content)
        reader = PdfReader(pdf_file)
        
        documents = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text.strip():
                chunk_metadata = metadata.copy()
                chunk_metadata['page'] = i + 1
                chunk_metadata['chunk_type'] = 'page'
                
                doc = Document(
                    content=text,
                    metadata=chunk_metadata,
                    doc_id=metadata['doc_id'],
                    chunk_id=f"{metadata['doc_id']}_page_{i+1}"
                )
                documents.append(doc)
        
        return documents
    
    def _parse_docx(self, content: bytes, metadata: Dict) -> List[Document]:
        """Parse DOCX content"""
        import io
        doc_file = io.BytesIO(content)
        doc = docx.Document(doc_file)
        
        full_text = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                full_text.append(paragraph.text)
        
        text = '\n'.join(full_text)
        return [Document(
            content=text,
            metadata=metadata,
            doc_id=metadata['doc_id']
        )]
    
    def _parse_text(self, content: bytes, metadata: Dict) -> List[Document]:
        """Parse plain text content"""
        text = content.decode('utf-8')
        return [Document(
            content=text,
            metadata=metadata,
            doc_id=metadata['doc_id']
        )]
    
    def _parse_json(self, content: bytes, metadata: Dict) -> List[Document]:
        """Parse JSON content"""
        data = json.loads(content.decode('utf-8'))
        documents = []
        
        if isinstance(data, list):
            for i, item in enumerate(data):
                chunk_metadata = metadata.copy()
                chunk_metadata['item_index'] = i
                
                # Handle different JSON structures
                if isinstance(item, dict):
                    # Try common QA patterns
                    if 'question' in item and 'answer' in item:
                        text = f"Q: {item['question']}\nA: {item['answer']}"
                        chunk_metadata['type'] = 'qa_pair'
                    else:
                        text = json.dumps(item, indent=2)
                else:
                    text = str(item)
                
                doc = Document(
                    content=text,
                    metadata=chunk_metadata,
                    doc_id=metadata['doc_id'],
                    chunk_id=f"{metadata['doc_id']}_item_{i}"
                )
                documents.append(doc)
        else:
            documents.append(Document(
                content=json.dumps(data, indent=2),
                metadata=metadata,
                doc_id=metadata['doc_id']
            ))
        
        return documents
    
    def _parse_csv(self, content: bytes, metadata: Dict) -> List[Document]:
        """Parse CSV content"""
        import io
        csv_file = io.StringIO(content.decode('utf-8'))
        reader = csv.DictReader(csv_file)
        
        documents = []
        for i, row in enumerate(reader):
            chunk_metadata = metadata.copy()
            chunk_metadata['row_index'] = i
            
            # Convert row to readable text
            text_parts = []
            for key, value in row.items():
                if value:
                    text_parts.append(f"{key}: {value}")
            
            text = '\n'.join(text_parts)
            
            doc = Document(
                content=text,
                metadata=chunk_metadata,
                doc_id=metadata['doc_id'],
                chunk_id=f"{metadata['doc_id']}_row_{i}"
            )
            documents.append(doc)
        
        return documents
    
    def _parse_html(self, content: bytes, metadata: Dict) -> List[Document]:
        """Parse HTML content"""
        soup = BeautifulSoup(content, 'html.parser')
        text = soup.get_text(separator='\n', strip=True)
        
        return [Document(
            content=text,
            metadata=metadata,
            doc_id=metadata['doc_id']
        )]

class DocumentChunker:
    """Splits documents into smaller chunks for better QA performance"""
    
    def __init__(self, chunk_size: int = 1000, overlap: int = 200):
        self.chunk_size = chunk_size
        self.overlap = overlap
    
    def chunk_document(self, document: Document) -> List[Document]:
        """Split document into chunks"""
        if len(document.content) <= self.chunk_size:
            return [document]
        
        chunks = []
        text = document.content
        
        for i in range(0, len(text), self.chunk_size - self.overlap):
            chunk_text = text[i:i + self.chunk_size]
            
            chunk_metadata = document.metadata.copy()
            chunk_metadata['chunk_index'] = len(chunks)
            chunk_metadata['chunk_start'] = i
            chunk_metadata['chunk_end'] = i + len(chunk_text)
            
            chunk = Document(
                content=chunk_text,
                metadata=chunk_metadata,
                doc_id=document.doc_id,
                chunk_id=f"{document.doc_id}_chunk_{len(chunks)}"
            )
            chunks.append(chunk)
        
        return chunks

class VectorDatabase:
    """ChromaDB implementation for vector storage"""
    
    def __init__(self, db_path: str = "./chroma_db", collection_name: str = "qa_collection"):
        self.client = chromadb.PersistentClient(path=db_path)
        self.collection_name = collection_name
        self.model = SentenceTransformer('all-MiniLM-L6-v2')
        
        # Create or get collection
        try:
            self.collection = self.client.get_collection(name=collection_name)
        except (ValueError, NotFoundError):
            self.collection = self.client.create_collection(
                name=collection_name,
                metadata={"hnsw:space": "cosine"}
            )
    
    def add_documents(self, documents: List[Document]) -> None:
        """Add documents to the vector database"""
        if not documents:
            return
        
        # Prepare data for ChromaDB
        ids = []
        texts = []
        metadatas = []
        
        for doc in documents:
            doc_id = doc.chunk_id if doc.chunk_id else doc.doc_id
            ids.append(doc_id)
            texts.append(doc.content)
            metadatas.append(doc.metadata)
        
        # Generate embeddings
        embeddings = self.model.encode(texts).tolist()
        
        # Add to ChromaDB
        self.collection.add(
            ids=ids,
            documents=texts,
            metadatas=metadatas,
            embeddings=embeddings
        )
    
    def search(self, query: str, n_results: int = 5, 
               filter_dict: Optional[Dict] = None) -> List[Dict]:
        """Search for similar documents"""
        query_embedding = self.model.encode([query]).tolist()
        
        results: QueryResult = self.collection.query(
            query_embeddings=query_embedding,
            n_results=n_results,
            where=filter_dict
        )
        # Type casting for safety
        ids: List[List[str]] = results['ids'] if isinstance(results['ids'], list) else list(results['ids'])
        documents: List[List[str]] = results['documents'] if isinstance(results['documents'], list) else list(results['documents'])
        metadatas: List[List[Metadata]] = results['metadatas'] if isinstance(results['metadatas'], list) else list(results['metadatas'])
        distances: List[List[float]] = results['distances'] if isinstance(results['distances'], list) else list(results['distances'])

        return [{
            'id': ids[0][i],
            'document': documents[0][i],
            'metadata': metadatas[0][i],
            'distance': distances[0][i]
        } for i in range(len(ids[0]))]
    
    def get_collection_stats(self) -> Dict:
        """Get collection statistics"""
        count = self.collection.count()
        return {
            'total_documents': count,
            'collection_name': self.collection_name
        }

class DocumentCollectionPipeline:
    """Main pipeline for processing and storing documents"""
    
    def __init__(self, db_path: str = "./chroma_db", collection_name: str = "qa_collection"):
        self.processor = FileProcessor()
        self.chunker = DocumentChunker()
        self.db = VectorDatabase(db_path, collection_name)
    
    def process_files(self, file_paths: List[str], source_type: str = "local") -> Dict:
        """Process multiple files and store in database"""
        total_docs = 0
        total_chunks = 0
        errors = []
        
        for file_path in file_paths:
            try:
                print(f"Processing: {file_path}")
                
                # Parse file
                documents = self.processor.process_file(file_path, source_type)
                
                # Chunk documents
                all_chunks = []
                for doc in documents:
                    chunks = self.chunker.chunk_document(doc)
                    all_chunks.extend(chunks)
                
                # Store in database
                self.db.add_documents(all_chunks)
                
                total_docs += len(documents)
                total_chunks += len(all_chunks)
                
                print(f"✓ Processed {len(documents)} documents, {len(all_chunks)} chunks")
                
            except Exception as e:
                error_msg = f"Error processing {file_path}: {str(e)}"
                errors.append(error_msg)
                print(f"✗ {error_msg}")
        
        return {
            'total_documents': total_docs,
            'total_chunks': total_chunks,
            'errors': errors,
            'collection_stats': self.db.get_collection_stats()
        }
    
    def search_collection(self, query: str, n_results: int = 5) -> List[Dict]:
        """Search the document collection"""
        return self.db.search(query, n_results)

# Example usage
if __name__ == "__main__":
    # Initialize pipeline
    pipeline = DocumentCollectionPipeline()
    
    # Process local files
    local_files = [
        "documents/Raju_Bhimani_Updated_Resume.docx",
        "documents/Raju_Bhimani_Updated_Resume.pdf",
        # "documents/data.csv"
    ]
    
    # Process S3 files (if AWS configured)
    s3_files = [
        # "my-bucket/documents/report.pdf",
        # "my-bucket/data/qa_pairs.json"
    ]
    
    # Process files
    print("Processing local files...")
    result = pipeline.process_files(local_files, "local")
    print(f"Results: {result}")
    
    # Search example
    print("\nSearching collection...")
    results = pipeline.search_collection("How to install the software?")
    for i, result in enumerate(results):
        print(f"{i+1}. {result['document'][:200]}...")
        print(f"   Source: {result['metadata']['source']}")
        print(f"   Distance: {result['distance']:.3f}\n")
